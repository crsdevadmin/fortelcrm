"""Read the weekly secondary-sales workbook that reps upload.

Layout (one sheet per month, e.g. "October"):

    row A:  |        |          |        | Stockist 1 (merged)          | Stockist 2 ...
    row B:  |        |          |        | Week 1 | Week 2 | ... | Total Month | ...
    row C:  | S.No   | Products | A-Rate | Sales | Closing | Sales | Closing ...
    rows :  |   1    | Amtrios  |  3416  |  17   |   10    | ...

Only the weekly Sales / Closing cells are read. "Total Month" and
"Total Region Sales" columns are ignored and recalculated here, because the
formulas in the circulated template are not reliable.

Secondary sale value = weekly sales quantity x rate. The rate comes from the
sheet's rate column; if that cell is blank the Product Master PTS is used.
"""

import io
import re

MONTH_NAMES = [
    "", "january", "february", "march", "april", "may", "june",
    "july", "august", "september", "october", "november", "december",
]

PRODUCT_HEADERS = {"products", "product", "productname", "productsname", "particulars", "itemname", "item"}
RATE_HEADERS = {"arate", "rate", "pts", "ptr", "purrate", "purchaserate", "salesrate"}
SALES_HEADERS = {"sales", "sale", "salesqty", "saleqty", "sqty"}
CLOSING_HEADERS = {"closing", "closingqty", "closingstock", "clstock", "clsqty", "stock"}

FORM_WORDS = {
    "cap": "cap", "caps": "cap", "capsule": "cap", "capsules": "cap", "softgel": "cap",
    "tab": "tab", "tabs": "tab", "tablet": "tab", "tablets": "tab",
    "inj": "inj", "injection": "inj", "vial": "inj",
    "susp": "susp", "suspension": "susp", "syrup": "susp", "oral": "susp",
    "spray": "spray", "lotion": "lotion", "cream": "cream", "gel": "gel",
    "mouthwash": "wash", "wash": "wash", "viscous": "viscous", "lp": "lp", "sachet": "sachet",
}


class SheetFormatError(ValueError):
    """The workbook does not look like the secondary-sales template."""


def _key(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _tokens(value):
    text = str(value or "").lower()
    text = re.sub(r"([a-z])(\d)", r"\1 \2", text)
    text = re.sub(r"(\d)([a-z])", r"\1 \2", text)
    return [token for token in re.split(r"[^a-z0-9.]+", text) if token and token != "."]


def _clean_name(value):
    return " ".join(str(value or "").replace("\n", " ").split()).strip(" ,")


def _number(value):
    """Return (number, problem). Blank cells are 0 with no problem."""
    if value is None:
        return 0.0, None
    if isinstance(value, bool):
        return 0.0, "not a number"
    if isinstance(value, (int, float)):
        return float(value), None
    text = str(value).strip().replace(",", "")
    if text in {"", "-", "--", "nil", "NIL", "Nil"}:
        return 0.0, None
    try:
        return float(text), None
    except ValueError:
        return 0.0, "not a number"


def _product_rate(product):
    return float(getattr(product, "price", None) or getattr(product, "rate", None) or 0)


def match_product(products, source_name, rate=0.0):
    """Match a sheet product name ("Caximeg240", "Amtrios-Inj") to the Product Master."""
    tokens = _tokens(source_name)
    if not tokens:
        return None
    brand = tokens[0]
    numbers = {t for t in tokens if re.fullmatch(r"\d+(?:\.\d+)?", t)}
    forms = {FORM_WORDS[t] for t in tokens if t in FORM_WORDS}
    candidates = []
    for product in products:
        name = getattr(product, "name", "") or ""
        product_tokens = _tokens(name)
        if not product_tokens or product_tokens[0] != brand:
            continue
        product_numbers = {t for t in product_tokens if re.fullmatch(r"\d+(?:\.\d+)?", t)}
        product_forms = {FORM_WORDS[t] for t in product_tokens if t in FORM_WORDS}
        number_score = 0
        if numbers:
            number_score = 0 if numbers & product_numbers else 1
        form_score = 0
        if forms:
            form_score = 0 if forms & product_forms else 1
        expected = _product_rate(product)
        rate_gap = abs(expected - rate) / rate if rate and expected else 1.0
        shared = len(set(tokens) & set(product_tokens))
        candidates.append((number_score, form_score, round(rate_gap, 4), -shared, len(name), product))
    if not candidates:
        return None
    best = min(candidates, key=lambda item: item[:-1])
    # A size written on the sheet (e.g. 240) that no product carries, with no
    # rate to confirm it, is too uncertain to map automatically.
    if best[0] and best[2] > 0.15 and len(candidates) > 1:
        return None
    return best[-1]


def _pick_sheet(workbook, month):
    names = workbook.sheetnames
    if not names:
        raise SheetFormatError("The workbook has no sheets")
    if month:
        full = MONTH_NAMES[month]
        short = full[:3]
        for name in names:
            if _key(name) == full:
                return workbook[name]
        for name in names:
            key = _key(name)
            if key.startswith(short) or full in key:
                return workbook[name]
        if len(names) == 1:
            return workbook[names[0]]
        raise SheetFormatError(
            f"No sheet named {full.title()} in this workbook. Sheets found: {', '.join(names)}"
        )
    return workbook[names[0]]


def _grid(sheet):
    """Sheet values as a list of rows with merged cells filled in."""
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    width = max((len(row) for row in rows), default=0)
    rows = [row + [None] * (width - len(row)) for row in rows]
    merged_ranges = sheet.merged_cells.ranges if getattr(sheet, "merged_cells", None) else []
    for merged in merged_ranges:
        top, left = merged.min_row - 1, merged.min_col - 1
        if top >= len(rows) or left >= width:
            continue
        value = rows[top][left]
        for r in range(top, min(merged.max_row, len(rows))):
            for c in range(left, min(merged.max_col, width)):
                rows[r][c] = value
    return rows


def _find_header(rows):
    for index, row in enumerate(rows[:40]):
        keys = [_key(cell) for cell in row]
        product_col = next((i for i, key in enumerate(keys) if key in PRODUCT_HEADERS), None)
        sales_cols = [i for i, key in enumerate(keys) if key in SALES_HEADERS]
        if product_col is not None and sales_cols:
            rate_col = next((i for i, key in enumerate(keys) if key in RATE_HEADERS), None)
            return index, product_col, rate_col
    raise SheetFormatError(
        "Could not find the header row. The sheet needs a 'Products' column and 'Sales' / 'Closing' columns under each week."
    )


def _week_number(label):
    key = _key(label)
    match = re.fullmatch(r"(?:week|wk|w)(\d)", key)
    if match:
        return int(match.group(1))
    return None


def _column_map(rows, header_index):
    """Return [(stockist, week, sales_col, closing_col)]."""
    header = rows[header_index]
    week_row = None
    for index in range(header_index - 1, max(-1, header_index - 4), -1):
        if any(_week_number(cell) for cell in rows[index]):
            week_row = index
            break
    if week_row is None:
        raise SheetFormatError("Could not find the 'Week 1 / Week 2 ...' row above the Sales / Closing headers.")
    stockist_row = week_row - 1 if week_row > 0 else None

    stockists = []
    current = None
    if stockist_row is not None:
        for cell in rows[stockist_row]:
            name = _clean_name(cell)
            if name:
                current = name
            stockists.append(current)
    else:
        stockists = [None] * len(header)

    columns = []
    used = set()
    for col, cell in enumerate(header):
        if _key(cell) not in SALES_HEADERS or col in used:
            continue
        week = _week_number(rows[week_row][col])
        stockist = stockists[col] if col < len(stockists) else None
        if not week or not stockist or "total" in _key(stockist):
            continue
        closing_col = None
        nxt = col + 1
        if nxt < len(header) and _key(header[nxt]) in CLOSING_HEADERS \
                and _week_number(rows[week_row][nxt]) == week \
                and (stockists[nxt] if nxt < len(stockists) else None) == stockist:
            closing_col = nxt
            used.add(nxt)
        columns.append((stockist, week, col, closing_col))
    if not columns:
        raise SheetFormatError("No stockist / week Sales columns were found under the header row.")
    return columns


def parse_secondary_sheet(raw, products, month=None, filename=""):
    """Parse the uploaded workbook.

    Returns a dict with ``lines`` (one per stockist x product x week with any
    sales or closing), ``stockists``, ``totals``, ``warnings`` and
    ``unmatched_products``.
    """
    lower = (filename or "").lower()
    if lower.endswith(".xls") and not lower.endswith(".xlsx"):
        raise SheetFormatError("Save the sheet as .xlsx (Excel Workbook) and upload again.")
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as exc:  # corrupt or not an xlsx
        raise SheetFormatError("This file could not be opened as an Excel .xlsx workbook.") from exc
    sheet = _pick_sheet(workbook, month)
    rows = _grid(sheet)
    header_index, product_col, rate_col = _find_header(rows)
    columns = _column_map(rows, header_index)

    warnings = []
    if rate_col is None:
        warnings.append("No rate column (A-Rate) found — Product Master rates were used for every product.")
    stockist_order = []
    for stockist, *_ in columns:
        if stockist not in stockist_order:
            stockist_order.append(stockist)

    lines = []
    unmatched = {}
    missing_rate = []
    bad_cells = []
    seen_products = {}
    blank_run = 0
    for row_index in range(header_index + 1, len(rows)):
        row = rows[row_index]
        source_name = _clean_name(row[product_col] if product_col < len(row) else None)
        if not source_name:
            blank_run += 1
            if blank_run >= 5:
                break
            continue
        blank_run = 0
        if _key(source_name).startswith("total") or _key(source_name) in {"grandtotal", "products"}:
            continue
        excel_row = row_index + 1

        sheet_rate, problem = _number(row[rate_col]) if rate_col is not None and rate_col < len(row) else (0.0, None)
        if problem:
            bad_cells.append(f"{source_name} rate")
            sheet_rate = 0.0
        product = match_product(products, source_name, sheet_rate)
        master_rate = _product_rate(product) if product else 0.0
        if sheet_rate > 0:
            rate, rate_source = sheet_rate, "sheet"
        elif master_rate > 0:
            rate, rate_source = master_rate, "product_master"
        else:
            rate, rate_source = 0.0, "missing"

        product_key = _key(source_name)
        if product_key in seen_products:
            warnings.append(f"'{source_name}' appears twice (rows {seen_products[product_key]} and {excel_row}); both rows were added together.")
        seen_products.setdefault(product_key, excel_row)

        row_has_data = False
        for stockist, week, sales_col, closing_col in columns:
            sales, sales_problem = _number(row[sales_col])
            closing, closing_problem = _number(row[closing_col]) if closing_col is not None else (0.0, None)
            if sales_problem:
                bad_cells.append(f"{source_name} · {stockist} · Week {week} sales")
            if closing_problem:
                bad_cells.append(f"{source_name} · {stockist} · Week {week} closing")
            if sales < 0 or closing < 0:
                warnings.append(f"Negative quantity for {source_name} at {stockist}, Week {week} was treated as 0.")
                sales, closing = max(sales, 0.0), max(closing, 0.0)
            if sales == 0 and closing == 0:
                continue
            row_has_data = True
            lines.append({
                "stockist": stockist,
                "week": min(week, 4),
                "source_product_name": source_name,
                "product_id": product.id if product else None,
                "product_name": product.name if product else None,
                "rate": round(rate, 2),
                "rate_source": rate_source,
                "sales_qty": sales,
                "sales_value": round(sales * rate, 2),
                "closing_qty": closing,
                "closing_value": round(closing * rate, 2),
            })
        if row_has_data and rate_source == "missing":
            missing_rate.append(source_name)
        if row_has_data and not product:
            unmatched[source_name] = unmatched.get(source_name, 0) + 1

    if any(week > 4 for _, week, _, _ in columns):
        warnings.append("Week 5 figures were added into Week 4.")
    if missing_rate:
        warnings.append(f"No rate on the sheet or in the Product Master for: {', '.join(missing_rate)}. Their value is counted as 0.")
    if bad_cells:
        shown = ", ".join(bad_cells[:6]) + (f" and {len(bad_cells) - 6} more" if len(bad_cells) > 6 else "")
        warnings.append(f"Text found where a number was expected (treated as 0): {shown}.")
    if unmatched:
        warnings.append(
            f"{len(unmatched)} product(s) are not in the Product Master: {', '.join(unmatched)}. "
            "They are included in secondary sales but not in product targets."
        )

    # Merge duplicates of the same stockist / product / week.
    merged = {}
    for line in lines:
        key = (line["stockist"], _key(line["source_product_name"]), line["week"])
        if key in merged:
            current = merged[key]
            for field in ("sales_qty", "sales_value", "closing_qty", "closing_value"):
                current[field] = round(current[field] + line[field], 2)
        else:
            merged[key] = dict(line)
    lines = list(merged.values())

    totals = {
        "sales_qty": round(sum(line["sales_qty"] for line in lines), 3),
        "sales_value": round(sum(line["sales_value"] for line in lines), 2),
    }
    latest_week = {}
    for line in lines:
        if line["closing_qty"] > 0:
            key = (line["stockist"], _key(line["source_product_name"]))
            if key not in latest_week or line["week"] > latest_week[key]["week"]:
                latest_week[key] = line
    totals["closing_qty"] = round(sum(line["closing_qty"] for line in latest_week.values()), 3)
    totals["closing_value"] = round(sum(line["closing_value"] for line in latest_week.values()), 2)

    return {
        "sheet_name": sheet.title,
        "stockists": stockist_order,
        "weeks_with_data": sorted({line["week"] for line in lines if line["sales_qty"] or line["closing_qty"]}),
        "lines": lines,
        "totals": totals,
        "warnings": warnings,
        "unmatched_products": list(unmatched),
    }
