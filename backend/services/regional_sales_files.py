"""Extract regional sales rows from Excel workbooks and report screenshots."""

import io
import re

from .regional_sales_pdf import _tally_product


def _key(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _number(value):
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _find_header(rows):
    name_keys = {"itemname", "productname", "particulars", "product", "item"}
    quantity_keys = {"sales", "salesqty", "saleqty", "sqty", "outwards", "outwardqty"}
    value_keys = {"salesval", "salesvalue", "salevalue", "salvalue", "outwardsvalue", "outwardvalue"}
    rate_keys = {"rate", "purrate", "purchaserate", "salesrate", "salerate"}
    for index, row in enumerate(rows):
        keys = [_key(cell) for cell in row]
        name = next((i for i, item in enumerate(keys) if item in name_keys), None)
        quantity = next((i for i, item in enumerate(keys) if item in quantity_keys), None)
        if name is None or quantity is None:
            continue
        value = next((i for i, item in enumerate(keys) if item in value_keys), None)
        rate = next((i for i, item in enumerate(keys) if item in rate_keys), None)
        return index, {"name": name, "quantity": quantity, "value": value, "rate": rate}
    return None, None


def _map_rows(rows, products):
    header_index, columns = _find_header(rows)
    if columns is None:
        return {"entries": [], "unmatched_rows": 0, "unmatched_items": [], "source_total": None, "matched_total": 0.0}
    entries = {}
    unmatched = 0
    unmatched_items = []
    source_total = 0.0
    matched_total = 0.0
    for row in rows[header_index + 1:]:
        if columns["name"] >= len(row):
            continue
        name = str(row[columns["name"]] or "").strip()
        if not name or _key(name) in {"total", "grandtotal"}:
            continue
        quantity = _number(row[columns["quantity"]] if columns["quantity"] < len(row) else None)
        value = _number(row[columns["value"]] if columns["value"] is not None and columns["value"] < len(row) else None)
        rate = _number(row[columns["rate"]] if columns["rate"] is not None and columns["rate"] < len(row) else None)
        if quantity <= 0:
            continue
        if rate <= 0 and value > 0:
            rate = value / quantity
        if rate <= 0:
            continue
        line_value = value if value > 0 else quantity * rate
        source_total += line_value
        product = _tally_product(products, name, rate)
        if not product:
            unmatched += 1
            unmatched_items.append({
                "source_name": name,
                "quantity": round(quantity, 3),
                "price": round(rate, 2),
                "value": round(line_value, 2),
            })
            continue
        matched_total += line_value
        current = entries.setdefault(product.id, {
            "product_id": product.id,
            "product_name": product.name,
            "source_name": name,
            "quantity": 0.0,
            "price": rate,
        })
        current["quantity"] += quantity
        current["price"] = rate
    result = list(entries.values())
    for row in result:
        row["quantity"] = round(row["quantity"], 3)
        row["price"] = round(row["price"], 2)
        row["value"] = round(row["quantity"] * row["price"], 2)
    return {
        "entries": result,
        "unmatched_rows": unmatched,
        "unmatched_items": unmatched_items,
        "source_total": round(source_total, 2),
        "matched_total": round(matched_total, 2),
    }


def extract_excel_rows(raw, products, filename=""):
    if (filename or "").lower().endswith(".xls"):
        import xlrd
        workbook = xlrd.open_workbook(file_contents=raw)
        rows = [sheet.row_values(index) for sheet in workbook.sheets() for index in range(sheet.nrows)]
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
        rows = [list(row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
    return _map_rows(rows, products)


def extract_image_rows(raw, products):
    from PIL import Image
    import pytesseract

    image = Image.open(io.BytesIO(raw))
    image.verify()
    image = Image.open(io.BytesIO(raw)).convert("RGB")
    if image.width * image.height > 25_000_000:
        raise ValueError("Image dimensions are too large")
    data = pytesseract.image_to_data(image, config="--psm 6", output_type=pytesseract.Output.DICT)
    lines = {}
    for index, text in enumerate(data["text"]):
        text = (text or "").strip()
        if not text:
            continue
        line_key = (data["block_num"][index], data["par_num"][index], data["line_num"][index])
        lines.setdefault(line_key, []).append({
            "text": text,
            "x": data["left"][index],
            "right": data["left"][index] + data["width"][index],
        })
    ordered = [sorted(words, key=lambda word: word["x"]) for words in lines.values()]
    header_index = next((index for index, words in enumerate(ordered)
                         if "itemname" in _key("".join(word["text"] for word in words))
                         and "sqty" in {_key(word["text"]) for word in words}), None)
    if header_index is None:
        return {"entries": [], "unmatched_rows": 0, "source_total": None}
    header = ordered[header_index]

    def header_x(*keys):
        keys = set(keys)
        matches = [word for word in header if _key(word["text"]) in keys]
        return matches[-1]["x"] if matches else None

    pack_x = header_x("pack")
    quantity_x = header_x("sqty", "salesqty")
    rate_x = header_x("purrate", "rate")
    sal_words = [word for word in header if _key(word["text"]) in {"salvalue", "sal", "salesvalue"}]
    value_x = sal_words[-1]["x"] if sal_words else None
    if None in (pack_x, quantity_x, value_x, rate_x):
        return {"entries": [], "unmatched_rows": 0, "source_total": None}

    def closest(words, x):
        numeric = [(abs(word["x"] - x), word["text"]) for word in words if re.fullmatch(r"-?[\d,.]+", word["text"])]
        return min(numeric)[1] if numeric and min(numeric)[0] < 45 else None

    rows = [["Item name", "SQty", "Sal Value", "Pur.Rate"]]
    for words in ordered[header_index + 1:]:
        name = " ".join(word["text"] for word in words if word["x"] < pack_x).strip()
        if not name or _key(name) in {"total", "grandtotal"}:
            continue
        rows.append([name, closest(words, quantity_x), closest(words, value_x), closest(words, rate_x)])
    return _map_rows(rows, products)
