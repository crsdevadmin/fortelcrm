import hashlib
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable


MAX_PRIMARY_SALES_ROWS = 20000


def normalize_stockist_name(value) -> str:
    return " ".join(str(value or "").strip().upper().split())


def _header_key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).strip().split())


def _number(value) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("₹", "")
    if not text:
        return 0.0
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    try:
        return float(text)
    except ValueError:
        return 0.0


def _iso_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return ""


REQUIRED_HEADERS = {
    "customer_name": {"customername", "stockistname", "distributorname"},
    "bill_number": {"billnumber", "invoicenumber", "invoiceno", "billno"},
    "bill_date": {"billdate", "invoicedate"},
    "product_name": {"productname", "itemname"},
    "quantity": {"quantity", "qty"},
    "net_amount": {"netamount", "netvalue"},
}

OPTIONAL_HEADERS = {
    "free_quantity": {"freequantity", "freeqty"},
    "rate": {"rate", "price"},
    "gross_amount": {"grossamountwithdiscount", "grossamount", "grossbeforediscount"},
    "tax_amount": {"taxamount", "gst"},
    "batch_number": {"batchdescription", "batchnumber", "batchno"},
    "gst_number": {"gstnumber", "gstin"},
}


def _column_map(header_row) -> dict[str, int]:
    normalized = [_header_key(value) for value in header_row]
    result = {}
    for field, aliases in {**REQUIRED_HEADERS, **OPTIONAL_HEADERS}.items():
        for index, key in enumerate(normalized):
            if key in aliases:
                result[field] = index
                break
    return result


def parse_primary_sales_rows(rows: Iterable[Iterable]) -> dict:
    materialized = [list(row) for row in rows]
    header_index = None
    columns = {}
    for index, row in enumerate(materialized[:30]):
        candidate = _column_map(row)
        if all(field in candidate for field in REQUIRED_HEADERS):
            header_index = index
            columns = candidate
            break
    if header_index is None:
        required = ", ".join(name.replace("_", " ").title() for name in REQUIRED_HEADERS)
        raise ValueError(f"Could not find the sales header row. Required columns: {required}")

    parsed = []
    skipped = 0
    occurrence_counts = {}

    def cell(row, field):
        index = columns.get(field)
        return row[index] if index is not None and index < len(row) else None

    for row in materialized[header_index + 1:]:
        stockist_name = _text(cell(row, "customer_name"))
        product_name = _text(cell(row, "product_name"))
        bill_number = _text(cell(row, "bill_number"))
        bill_date = _iso_date(cell(row, "bill_date"))
        if not any(_text(value) for value in row):
            continue
        if not stockist_name or stockist_name.lower() in {"grand total", "total"}:
            skipped += 1
            continue
        if not product_name or not bill_number or not bill_date:
            skipped += 1
            continue

        batch_number = _text(cell(row, "batch_number"))
        rate = _number(cell(row, "rate"))
        base_key = "|".join((
            normalize_stockist_name(stockist_name),
            bill_number.upper(),
            bill_date,
            product_name.upper(),
            batch_number.upper(),
            f"{rate:.4f}",
        ))
        occurrence_counts[base_key] = occurrence_counts.get(base_key, 0) + 1
        source_key = hashlib.sha256(
            f"{base_key}|{occurrence_counts[base_key]}".encode("utf-8")
        ).hexdigest()

        parsed.append({
            "source_key": source_key,
            "stockist_name": stockist_name,
            "normalized_stockist_name": normalize_stockist_name(stockist_name),
            "bill_number": bill_number,
            "bill_date": bill_date,
            "product_name": product_name,
            "batch_number": batch_number,
            "quantity": _number(cell(row, "quantity")),
            "free_quantity": _number(cell(row, "free_quantity")),
            "rate": rate,
            "gross_amount": _number(cell(row, "gross_amount")),
            "net_amount": _number(cell(row, "net_amount")),
            "tax_amount": _number(cell(row, "tax_amount")),
            "gst_number": _text(cell(row, "gst_number")),
        })
        if len(parsed) > MAX_PRIMARY_SALES_ROWS:
            raise ValueError(f"The workbook contains more than {MAX_PRIMARY_SALES_ROWS:,} sales rows")

    if not parsed:
        raise ValueError("No valid primary-sales rows were found in the workbook")
    return {"rows": parsed, "skipped_count": skipped, "header_row": header_index + 1}


def _read_xls(content: bytes):
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("Legacy .xls support is not installed on the server") from exc
    workbook = xlrd.open_workbook(file_contents=content)
    sheet = workbook.sheet_by_index(0)
    rows = []
    for row_index in range(sheet.nrows):
        values = []
        for column_index in range(sheet.ncols):
            cell = sheet.cell(row_index, column_index)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, workbook.datemode)
            values.append(value)
        rows.append(values)
    return rows


def _read_xlsx(content: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Modern Excel support is not installed on the server") from exc
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def parse_primary_sales_workbook(content: bytes, filename: str) -> dict:
    suffix = Path(filename or "").suffix.lower()
    try:
        if suffix == ".xls":
            rows = _read_xls(content)
        elif suffix == ".xlsx":
            rows = _read_xlsx(content)
        else:
            raise ValueError("Please upload an Excel .xls or .xlsx file")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The Excel file could not be read. Please upload a valid .xls or .xlsx workbook") from exc
    parsed = parse_primary_sales_rows(rows)
    parsed["file_checksum"] = hashlib.sha256(content).hexdigest()
    return parsed


SEED_STOCKISTS = (
    ("CBE HEXACARE", "Tamil Nadu", "Coimbatore 1"),
    ("CONNECT PHARMA DIST AND SUPPLIERS", "Tamil Nadu", "Coimbatore 1"),
    ("CLASSIC PHAR", "Tamil Nadu", "Coimbatore 1"),
    ("DR.N.RAJESH KAR", "Tamil Nadu", "Chennai"),
    ("EAZYMED TECHNOLOGIES P LTD", "Tamil Nadu", "Chennai"),
    ("ERODE JEYAM LIFE CARE", "Tamil Nadu", "Coimbatore 1"),
    ("EVERGREEN ENTERPRISES", "Tamil Nadu", "Chennai"),
    ("GURULAKSHMI PHARMA", "Tamil Nadu", "Chennai"),
    ("HARISH PHARMA", "Tamil Nadu", "Madurai"),
    ("HERMES HEALTH MART", "Tamil Nadu", "Coimbatore 1"),
    ("HEXACARE PHARMACEUTICALS PVT LTD", "Tamil Nadu", "Chennai"),
    ("JAYAM LIFE CARE", "Tamil Nadu", "Madurai"),
    ("JEI SPECIALITY DRUGS", "Tamil Nadu", "Coimbatore 1"),
    ("JH HEALTH CARE", "Tamil Nadu", "Madurai"),
    ("KALPANA", "Tamil Nadu", "Madurai"),
    ("KAMALAM MEDICAL CORPORATION", "Tamil Nadu", "Chennai"),
    ("KAVIN MEDICALS", "Tamil Nadu", "Chennai"),
    ("LIFE CARE DRUGS", "Tamil Nadu", "Chennai"),
    ("LIFECARE PHARMA PVT LTD", "Tamil Nadu", "Chennai"),
    ("MAHIMA CANCER FOUNTATION", "Tamil Nadu", "Chennai"),
    ("MARS MEDICALS&DISTRIBUTORS", "Tamil Nadu", "Chennai"),
    ("MAX LIFE SCIENCES", "Tamil Nadu", "Madurai"),
    ("MEDICINE HOUSE COIMBATORE", "Tamil Nadu", "Coimbatore 1"),
    ("MEDICINE HOUSE SALEM", "Tamil Nadu", "Coimbatore 1"),
    ("MEDICINE POINT", "Tamil Nadu", "Chennai"),
    ("MEDIHAUXE INTERNATIONAL INDIA ( P) LTD", "Tamil Nadu", "Chennai"),
    ("NEST PHARMA", "Tamil Nadu", "Chennai"),
    ("NEST BIO PHARMA", "Tamil Nadu", "Madurai"),
    ("NISHITA MEDICARE (P) LTD", "Tamil Nadu", "Coimbatore 1"),
    ("OPTIVAL HEALTH SOLUTIONS(P) LTD", "Tamil Nadu", "Chennai"),
    ("PLACENTA PHARMA", "Tamil Nadu", "Chennai"),
    ("PURANI HOSPITAL SUPPLIES (p) LTD", "Tamil Nadu", "Coimbatore 1"),
    ("SOORYA TRADING COMPANY", "Tamil Nadu", "Madurai"),
    ("SREE SARU PHARMA", "Tamil Nadu", "Coimbatore 1"),
    ("SREE VELMURUGAN PHARMA", "Tamil Nadu", "Madurai"),
    ("UNIQUE PHARMA", "Tamil Nadu", "Chennai"),
    ("V.S.BIOTECH UNIT OF V.S. HOSPITAL PVT LTD", "Tamil Nadu", "Chennai"),
    ("VEDA PHARMA", "Tamil Nadu", "Chennai"),
    ("NEXUS BIOCARE", "Tamil Nadu", "Unassigned"),
)


def ensure_seed_stockists(db):
    from ..models.models import Stockist

    existing = {row.normalized_name: row for row in db.query(Stockist).all()}
    for name, region, territory in SEED_STOCKISTS:
        normalized = normalize_stockist_name(name)
        if normalized not in existing:
            row = Stockist(
                name=name,
                normalized_name=normalized,
                region=region,
                territory=territory,
            )
            db.add(row)
            existing[normalized] = row
    db.flush()
    return existing


def infer_stockist_location(gst_number: str, stockist_name: str) -> tuple[str, str]:
    gst = str(gst_number or "").strip()
    name = normalize_stockist_name(stockist_name)
    if gst.startswith("36") or name.endswith("-HYD") or "HYDERABAD" in name:
        return "Telangana", "Hyderabad"
    if gst.startswith("32"):
        return "Kerala", "Cochin"
    if gst.startswith("33"):
        return "Tamil Nadu", "Unassigned"
    return "Unassigned", "Unassigned"
