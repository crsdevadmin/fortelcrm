import io
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from backend.services.secondary_sales_sheet import (
    SheetFormatError, match_product, parse_secondary_sheet,
)


def product(product_id, name, price):
    return SimpleNamespace(id=product_id, name=name, price=price, rate=price)


PRODUCTS = [
    product(1, "Amtrios Softgel Cap 30", 3416),
    product(2, "Amtrios Injection", 800),
    product(3, "Caximeg Oral Susp 60ml", 805),
    product(4, "Caximeg Oral Susp 240ml", 2956),
    product(5, "Zyora Mouthwash 200ml", 162),
    product(6, "Zyora Mouthwash 500ml", 369),
]


def workbook_bytes(sheet_name="October", stockists=("Medicine House CBE", "Nest"), rows=(), rate=True):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    first = 4 if rate else 3
    col = first
    for stockist in stockists:
        ws.cell(1, col, stockist)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 9)
        for index, label in enumerate(["Week 1", "Week 2", "Week 3", "Week 4", "Total Month"]):
            ws.cell(2, col + index * 2, label)
            ws.merge_cells(start_row=2, start_column=col + index * 2, end_row=2, end_column=col + index * 2 + 1)
            ws.cell(3, col + index * 2, "Sales")
            ws.cell(3, col + index * 2 + 1, "Closing")
        col += 10
    ws.cell(1, col, "Total Region Sales")
    ws.cell(3, col, "Sales")
    ws.cell(3, col + 1, "Closing")
    ws.cell(3, 1, "S.No")
    ws.cell(3, 2, "Products ")
    if rate:
        ws.cell(3, 3, "A-Rate")
    for offset, values in enumerate(rows):
        r = 4 + offset
        ws.cell(r, 1, offset + 1)
        for c, value in values.items():
            ws.cell(r, c, value)
        ws.cell(r, col, 999)  # unreliable template total must be ignored
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_reads_each_stockist_week_and_values_sales_at_sheet_rate():
    raw = workbook_bytes(rows=[
        {2: "Zyora 200ml", 3: 160, 4: 10, 5: 30, 6: 4, 7: 26, 14: 2, 15: 9},
    ])
    result = parse_secondary_sheet(raw, PRODUCTS, month=10, filename="team.xlsx")
    assert result["sheet_name"] == "October"
    assert result["stockists"] == ["Medicine House CBE", "Nest"]
    by_key = {(line["stockist"], line["week"]): line for line in result["lines"]}
    assert by_key[("Medicine House CBE", 1)]["sales_value"] == 1600
    assert by_key[("Medicine House CBE", 1)]["closing_qty"] == 30
    assert by_key[("Medicine House CBE", 2)]["sales_qty"] == 4
    assert by_key[("Nest", 1)]["sales_qty"] == 2
    assert by_key[("Nest", 1)]["product_id"] == 5
    assert result["totals"]["sales_qty"] == 16
    assert result["totals"]["sales_value"] == 2560
    # Closing uses the latest week per stockist/product, not the sum of weeks.
    assert result["totals"]["closing_qty"] == 26 + 9


def test_blank_rate_falls_back_to_product_master():
    raw = workbook_bytes(rows=[{2: "Amtrios Caps", 4: 17, 5: 10}])
    line = parse_secondary_sheet(raw, PRODUCTS, month=10)["lines"][0]
    assert line["product_id"] == 1
    assert line["rate"] == 3416
    assert line["rate_source"] == "product_master"
    assert line["sales_value"] == 17 * 3416


def test_unknown_product_is_kept_with_sheet_rate_and_warned():
    raw = workbook_bytes(rows=[{2: "Lycotrum", 3: 1363, 4: 2}])
    result = parse_secondary_sheet(raw, PRODUCTS, month=10)
    assert result["lines"][0]["product_id"] is None
    assert result["lines"][0]["sales_value"] == 2726
    assert result["unmatched_products"] == ["Lycotrum"]


def test_picks_the_sheet_for_the_selected_month():
    wb_bytes = workbook_bytes(sheet_name="November", rows=[{2: "Zyora 500ml", 3: 369, 4: 1}])
    result = parse_secondary_sheet(wb_bytes, PRODUCTS, month=11)
    assert result["sheet_name"] == "November"
    with pytest.raises(SheetFormatError):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(wb_bytes))
        wb.create_sheet("December")
        buffer = io.BytesIO()
        wb.save(buffer)
        parse_secondary_sheet(buffer.getvalue(), PRODUCTS, month=10)


def test_size_in_name_selects_the_right_pack():
    assert match_product(PRODUCTS, "Caximeg240", 2956).id == 4
    assert match_product(PRODUCTS, "Caximeg60", 805).id == 3
    assert match_product(PRODUCTS, "Amtrios-Inj", 0).id == 2
    assert match_product(PRODUCTS, "Zyora 500ml", 0).id == 6


def test_text_in_quantity_cell_is_reported():
    raw = workbook_bytes(rows=[{2: "Zyora 200ml", 3: 162, 4: "ten", 6: 3}])
    result = parse_secondary_sheet(raw, PRODUCTS, month=10)
    assert [line["sales_qty"] for line in result["lines"]] == [3]
    assert any("Text found" in warning for warning in result["warnings"])
